"""
Fyller i SVOA:s kalkylmall "Mall Kalkyl inriktningsbeslut LEDN" ur ledning_analys.py – en ifylld mall per område
och alternativ (alt. schakt: alla ledningar läggs om; alt. schaktfritt: infodring, rörspräckning där infodring inte
går och omläggning där inget av dem går).

Mallen redigeras direkt i filens XML (cellvärden i bladen), så att rullistor, bilder, kommentarer och formler i
mallen följer med oförändrade – openpyxl skulle ta bort rullistorna och bilderna. Det som fylls i:

  Uppgifter           projektnamn, alternativets nummer, beskrivning, meter per delsträcka (en delsträcka per
                      prisrad: rörspräckning ledning 1, infodring ledning 1, helschakt byggsträcka), servismeter,
                      antal servis- och ventilgropar
  Entreprenadkalkyl   prisraden (rullistans text) per delsträcka, införings- och maskingropar, servistillägg,
                      djuptillägg, slagen spont och provisorier (fritextrader i tabell D) och trafiktillägg
                      (tabell E, områdets procent)
  Rimlighetskontroll  byggherrekostnaderna = procentsatserna i fliken Alternativ (mallens BP2 hämtar annars
                      timkalkylen, som är tom)
  Kända risker        en fritextrad per fas med KANDA_RISKER_ANDEL av byggherrekostnaderna (+ entreprenaden)

Bara mallens poster används: hattar räknas som servisinfodring per meter, och ny brunn/framschaktning ingår inte.
Ett pris som saknas i mallens INDATA-flikar (t.ex. infodring Ø 1000) skrivs som fast pris på raden.
"""
import copy
import os
import re
import zipfile
from collections import defaultdict

from lxml import etree

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PNS = "http://schemas.openxmlformats.org/package/2006/relationships"
CNS = "http://schemas.openxmlformats.org/package/2006/content-types"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

MALL_MONSTER = re.compile(r"kalkyl.*inriktningsbeslut.*\.xlsx$", re.I)
OMRADE_FALT = "OMRADE"          # områdets namn på ledningen när Ledningsinventering inte satt 'omrade' (Exportera område)
UTAN_OMRADE = "Utan område"
ALTERNATIV = (("schakt", 1, "Alt. schakt – alla ledningar läggs om i helschakt"),
              ("schaktfritt", 2, "Alt. schaktfritt – infodring, rörspräckning där infodring inte går, omläggning "
                                 "där inget av dem går"))
DELSTRACKOR = 25                 # rader per tabell i mallen (Uppgifter rad 39–63)

# Mallens celler (mallversion 2026-08-01)
UPPG, ENTR, RIML, RISK = "Uppgifter ", "Entreprenadkalkyl", "Rimlighetskontroll", "Kända risker"
RAD0_UPPG = 39
RAD0_RORSP, RAD0_INFOD, RAD0_SCHAKT = 27, 63, 104
KOL_UPPG = {"rorsp": "E", "infod": "M", "schakt": "W"}
SERVIS_INFOD = {"V": "S", "S": "T", "K": "T", "D": "U"}       # Uppgifter: dricks-, spill-, dagvattenservis infodring
SERVIS_HEL = {"dricks": "AD", "komplett": "AE"}
GROP_SERVIS_UPPG = "K"
ENTR_INFORINGSGROP, ENTR_MASKINGROP = "D23", "D24"
ENTR_TILLAGG = {"bergschakt": "D97", "spontslade": "D98"}
ENTR_DJUP = ("D101", "D102")      # ≤ Ø 600, > Ø 600
ENTR_SPONT_RAD, ENTR_PROV_RAD = 146, 147   # fritextrader i tabell D
ENTR_TRAFIK_RAD = 153                      # tabell E innerstadstillägg nivå 1 (F = 1, G = procent)
ENTR_SUMMA_FORE_E = "(H55+H91+H132+H149)"
RIML_GRUPPER = ("24", "30", "37", "46", "65", "70", "75", "83")
RISK_PLANERA_RAD, RISK_GENOMFORA_RAD = 51, 103


# ============================================================================
# Redigering av xlsx på XML-nivå
# ============================================================================

class Formel(str):
    """Cellvärde som är en formel (utan '=')."""


def _kol_index(kol: str) -> int:
    n = 0
    for ch in kol:
        n = n * 26 + ord(ch) - 64
    return n


def _dela_ref(ref: str):
    m = re.match(r"([A-Z]+)(\d+)$", ref)
    return m.group(1), int(m.group(2))


class Arbetsbok:
    """Läser en xlsx som zip och ändrar cellvärden i bladens XML; allt annat skrivs tillbaka orört."""

    def __init__(self, path: str):
        with zipfile.ZipFile(path) as z:
            self.ordning = z.infolist()
            self.filer = {i.filename: z.read(i.filename) for i in self.ordning}
        wb = etree.fromstring(self.filer["xl/workbook.xml"])
        rels = etree.fromstring(self.filer["xl/_rels/workbook.xml.rels"])
        mal = {r.get("Id"): r.get("Target") for r in rels}
        self.blad = {}
        for sh in wb.iter(f"{{{NS}}}sheet"):
            t = mal[sh.get(f"{{{RNS}}}id")]
            self.blad[sh.get("name")] = t.lstrip("/") if t.startswith("/") else "xl/" + t
        self._xml = {}

    def _rot(self, blad: str):
        path = self.blad[blad]
        if path not in self._xml:
            self._xml[path] = etree.fromstring(self.filer[path])
        return self._xml[path]

    def _cell(self, blad: str, ref: str):
        rot = self._rot(blad)
        data = rot.find(f"{{{NS}}}sheetData")
        kol, nr = _dela_ref(ref)
        rad = None
        for r in data.iterfind(f"{{{NS}}}row"):
            rnr = int(r.get("r"))
            if rnr == nr:
                rad = r
                break
            if rnr > nr:
                rad = etree.Element(f"{{{NS}}}row", r=str(nr))
                r.addprevious(rad)
                break
        if rad is None:
            rad = etree.SubElement(data, f"{{{NS}}}row", r=str(nr))
        rad.attrib.pop("spans", None)
        for c in rad.iterfind(f"{{{NS}}}c"):
            ckol, _ = _dela_ref(c.get("r"))
            if ckol == kol:
                return c
            if _kol_index(ckol) > _kol_index(kol):
                ny = etree.Element(f"{{{NS}}}c", r=ref)
                c.addprevious(ny)
                return ny
        return etree.SubElement(rad, f"{{{NS}}}c", r=ref)

    def satt(self, blad: str, ref: str, varde) -> None:
        c = self._cell(blad, ref)
        for barn in list(c):
            c.remove(barn)
        c.attrib.pop("t", None)
        if varde is None:
            return
        if isinstance(varde, Formel):
            etree.SubElement(c, f"{{{NS}}}f").text = str(varde).lstrip("=")
        elif isinstance(varde, bool):
            c.set("t", "b")
            etree.SubElement(c, f"{{{NS}}}v").text = "1" if varde else "0"
        elif isinstance(varde, (int, float)):
            etree.SubElement(c, f"{{{NS}}}v").text = repr(round(float(varde), 6)) if isinstance(varde, float) \
                else str(varde)
        else:
            c.set("t", "inlineStr")
            t = etree.SubElement(etree.SubElement(c, f"{{{NS}}}is"), f"{{{NS}}}t")
            t.text = str(varde)
            t.set(XML_SPACE, "preserve")

    def _rensa_berakningar(self) -> None:
        """Tar bort cachade formelresultat (allt räknas när filen öppnas) utom i formler med externa länkar,
        och calcChain.xml; sätter fullCalcOnLoad."""
        for path in sorted(set(self.blad.values())):
            if path not in self.filer:
                continue
            rot = self._rot_path(path)
            delade = {}
            for f in rot.iter(f"{{{NS}}}f"):
                if f.get("t") == "shared" and f.text:
                    delade[f.get("si")] = f.text
            for c in rot.iter(f"{{{NS}}}c"):
                f = c.find(f"{{{NS}}}f")
                if f is None:
                    continue
                text = f.text or delade.get(f.get("si"), "")
                if "[" in text:              # extern länk – behåll värdet
                    continue
                v = c.find(f"{{{NS}}}v")
                if v is not None:
                    c.remove(v)
                c.attrib.pop("t", None)
        self.filer.pop("xl/calcChain.xml", None)
        rels = etree.fromstring(self.filer["xl/_rels/workbook.xml.rels"])
        for r in list(rels):
            if r.get("Target", "").endswith("calcChain.xml"):
                rels.remove(r)
        self.filer["xl/_rels/workbook.xml.rels"] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8",
                                                                  standalone=True)
        ct = etree.fromstring(self.filer["[Content_Types].xml"])
        for o in list(ct):
            if o.get("PartName", "").endswith("calcChain.xml"):
                ct.remove(o)
        self.filer["[Content_Types].xml"] = etree.tostring(ct, xml_declaration=True, encoding="UTF-8", standalone=True)
        wb = etree.fromstring(self.filer["xl/workbook.xml"])
        calc = wb.find(f"{{{NS}}}calcPr")
        if calc is None:
            calc = etree.SubElement(wb, f"{{{NS}}}calcPr")
        calc.set("fullCalcOnLoad", "1")
        calc.attrib.pop("calcCompleted", None)
        self.filer["xl/workbook.xml"] = etree.tostring(wb, xml_declaration=True, encoding="UTF-8", standalone=True)

    def _rot_path(self, path: str):
        if path not in self._xml:
            self._xml[path] = etree.fromstring(self.filer[path])
        return self._xml[path]

    def spara(self, ut: str) -> None:
        self._rensa_berakningar()
        for path, rot in self._xml.items():
            self.filer[path] = etree.tostring(rot, xml_declaration=True, encoding="UTF-8", standalone=True)
        tmp = ut + ".tmp"
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
            for info in self.ordning:
                if info.filename in self.filer:
                    ny = copy.copy(info)
                    ny.compress_type = zipfile.ZIP_DEFLATED
                    z.writestr(ny, self.filer[info.filename])
        os.replace(tmp, ut)


# ============================================================================
# Mallens prislistor (INDATA-flikarna)
# ============================================================================

def las_prislistor(mall: str) -> dict:
    """{'rorsp': {pris: text}, 'infod': {pris: text}, 'schakt': {(schakttyp, pris): text}, 'djup': (pris, pris),
    'servis_dricks': pris} ur mallen (första raden med priset vinner, som VLOOKUP)."""
    from openpyxl import load_workbook
    wb = load_workbook(mall, read_only=True, data_only=True)
    ut = {"rorsp": {}, "infod": {}, "schakt": {}}

    def rader(blad):
        for r in wb[blad].iter_rows(min_row=1, max_col=2, values_only=True):
            yield r[0], r[1]
    for blad, nyckel in (("INDATA RÖRSPRÄCKNING", "rorsp"), ("INDATA INFODRING", "infod")):
        for a, b in rader(blad):
            if isinstance(a, str) and isinstance(b, (int, float)) and b > 0:
                ut[nyckel].setdefault(float(b), a)
    sektion = None
    for a, b in rader("INDATA SCHAKT"):
        if isinstance(a, str) and b is None:
            h = a.upper()
            sektion = ("spontslade" if "SPONT" in h else "bergschakt" if "BERG" in h else "kc" if "KC" in h
                       else "ledningschakt" if "SCHAKT" in h and "SERVIS" not in h and "HELSCHAKT" not in h else None)
            continue
        if sektion and isinstance(a, str) and isinstance(b, (int, float)) and b > 0:
            ut["schakt"].setdefault((sektion, float(b)), a)
    ws = wb[ENTR]
    ut["djup"] = (ws["F101"].value, ws["F102"].value)
    ut["servis_dricks"] = ws["F95"].value
    wb.close()
    return ut


def hitta_mall(mappar) -> str | None:
    """Första filen som ser ut som kalkylmallen i någon av mapparna (inte Excels låsfiler)."""
    for m in mappar:
        if not m or not os.path.isdir(m):
            continue
        for f in sorted(os.listdir(m)):
            if MALL_MONSTER.search(f) and not f.startswith("~$"):
                return os.path.join(m, f)
    return None


# ============================================================================
# Underlag per område och alternativ
# ============================================================================

def omrade(l: dict) -> str:
    if l.get("omrade"):
        return str(l["omrade"])
    v = (l.get("attribut") or {}).get(OMRADE_FALT)
    if v not in (None, ""):
        return str(v).split(";")[0].strip() or UTAN_OMRADE
    return UTAN_OMRADE


class Underlag:
    def __init__(self, namn: str, alternativ: str):
        self.namn, self.alternativ = namn, alternativ
        self.rorsp = defaultdict(float)          # pris -> m
        self.infod = defaultdict(float)
        self.schakt = defaultdict(float)         # (schakttyp, pris) -> m grav
        self.schakt_text = {}
        self.djup = defaultdict(float)           # pris -> m (antal tillägg × meter)
        self.servis_infod = defaultdict(float)   # kolumn i Uppgifter -> m
        self.servis_hel = defaultdict(float)
        self.servis_tillagg = defaultdict(float)
        self.inforingsgropar = 0
        self.servis_ventilgropar = 0
        self.spont_m2 = 0.0
        self.ledningar = 0
        self.langd = 0.0
        self.metod_m = defaultdict(float)
        self.utan_pris = []
        self.kostnad = 0.0                       # mallens poster utan påslag (kontroll)
        self.trafik = []
        self.provisorier = []


def _servis_meter(serv, u: Underlag, kost, la):
    summa = 0.0
    for x in serv:
        m = la._servis_m(x)
        sy = x.get("system") or "S"
        u.servis_infod[SERVIS_INFOD.get(sy, "T")] += m
        summa += (la.kr(kost, "servis_infodring", system=sy) or 0.0) * m
    return summa


def _omlagg(o: dict, u: Underlag, dricks_pris, la, kost) -> float:
    summa = 0.0
    for p in o.get("poster") or []:
        if p["post"] == "schakt":
            nyckel = (p["typ"], float(p["pris"]))
            u.schakt[nyckel] += p["m"]
            u.schakt_text.setdefault(nyckel, f"{la.SCHAKTTYPER.get(p['typ'], p['typ'])}, {p['n']} ledning"
                                             f"{'ar' if p['n'] > 1 else ''}, Ø {p['dmax']:g} (kostnader_inventering.csv)")
            summa += p["pris"] * p["m"]
        elif p["post"] == "djuptillagg":
            u.djup[float(p["pris"])] += p["m"]
            summa += p["pris"] * p["m"]
    u.spont_m2 += o.get("spont_m2") or 0.0
    summa += o.get("spont") or 0.0
    for s in o.get("servis_poster") or []:
        hel = "dricks" if s.get("system") == "V" and dricks_pris and abs(s["pris"] - dricks_pris) < 0.5 \
            else "komplett"
        u.servis_hel[hel] += s["m"]
        summa += s["pris"] * s["m"]
        if s.get("tillagg_typ") in ENTR_TILLAGG:
            u.servis_tillagg[s["tillagg_typ"]] += s["m"]
            summa += s["tillagg"] * s["m"]
    return summa


def underlag(data: dict, ledningar: list[dict], kost, la, prislistor) -> dict:
    """{(område, alternativ): Underlag}."""
    serv_per = defaultdict(list)
    for x in data.get("_serviser_kopplade") or []:
        if x.get("ledning"):
            serv_per[x["ledning"]].append(x)
    ut = {}
    for l in ledningar:
        omr = omrade(l)
        serv = serv_per.get(l.get("id"), [])
        for alt, _, _ in ALTERNATIV:
            u = ut.setdefault((omr, alt), Underlag(omr, alt))
            if alt == "schakt":
                metod, d = "omläggning", l.get("_omlagg")
            else:
                metod, d = l.get("_alt_metod"), l.get("_alt")
            if not d:
                u.utan_pris.append(str(l.get("id")))
                continue
            L = l.get("langd_m") or 0.0
            u.ledningar += 1
            u.langd += L
            u.metod_m[metod] += L
            u.trafik.append(la.trafikpaslag(l))
            u.provisorier.append(la.provisorier(l))
            if metod == "strumpa":
                u.infod[float(d["per_meter"])] += L
                u.kostnad += d["infodring"] + _servis_meter(serv, u, kost, la)
            elif metod == "rörspräckning":
                u.rorsp[float(d["per_meter"])] += L
                u.inforingsgropar += d.get("antal_gropar") or 0
                u.servis_ventilgropar += len(serv) + (d.get("antal_ventilgropar") or 0)
                u.kostnad += d["sprackning"] + d["gropar"] + d["servisgropar"] + d["ventilgropar"] + \
                    _servis_meter(serv, u, kost, la)
            else:
                u.kostnad += _omlagg(d, u, prislistor.get("servis_dricks"), la, kost)
    return ut


def _vanligast(varden):
    v = [x for x in varden if x is not None]
    return max(set(v), key=v.count) if v else None


# ============================================================================
# Ifyllning
# ============================================================================

def _rader(poster: dict, lista: dict, text_for) -> list[tuple[str, float | None, float]]:
    """[(text i rullistan, fast pris eller None, meter)] sorterat på pris, högst DELSTRACKOR rader; prisrader som
    inte finns i mallen får fast pris, och fler rader än mallen rymmer slås ihop i sista raden."""
    rader = []
    for nyckel, m in sorted(poster.items(), key=lambda x: (str(x[0][0]), x[0][1]) if isinstance(x[0], tuple)
                            else ("", x[0])):
        if m <= 0:
            continue
        pris = nyckel[1] if isinstance(nyckel, tuple) else nyckel
        text = lista.get(nyckel)
        rader.append((text, None, m) if text else (text_for(nyckel), pris, m))
    if len(rader) > DELSTRACKOR:
        resten = rader[DELSTRACKOR - 1:]
        m = sum(r[2] for r in resten)
        kr = sum((r[1] if r[1] is not None else _pris(r, lista)) * r[2] for r in resten)
        rader = rader[:DELSTRACKOR - 1] + [(f"Övriga prisrader ({len(resten)} st, medelpris)", kr / m if m else 0, m)]
    return rader


def _pris(rad, lista):
    for nyckel, text in lista.items():
        if text == rad[0]:
            return nyckel[1] if isinstance(nyckel, tuple) else nyckel
    return 0.0


def fyll(mall: str, ut: str, u: Underlag, alt_nr: int, alt_text: str, prislistor: dict, kost, la,
         datum: str) -> None:
    wb = Arbetsbok(mall)
    tabeller = (
        ("rorsp", RAD0_RORSP, u.rorsp, prislistor["rorsp"], lambda p: f"Rörspräckning {p:g} kr/m (kostnader_inventering.csv)"),
        ("infod", RAD0_INFOD, u.infod, prislistor["infod"], lambda p: f"Infodring {p:g} kr/m (kostnader_inventering.csv)"),
        ("schakt", RAD0_SCHAKT, u.schakt, prislistor["schakt"], lambda k: u.schakt_text.get(k, f"Helschakt {k[1]:g} kr/m")),
    )
    for nyckel, rad0, poster, lista, text_for in tabeller:
        for i, (text, pris, m) in enumerate(_rader(poster, lista, text_for)):
            wb.satt(ENTR, f"B{rad0 + i}", text)
            if pris is not None:
                wb.satt(ENTR, f"F{rad0 + i}", round(pris, 2))
            wb.satt(UPPG, f"{KOL_UPPG[nyckel]}{RAD0_UPPG + i}", round(m, 2))
    for kol, m in u.servis_infod.items():
        wb.satt(UPPG, f"{kol}{RAD0_UPPG}", round(m, 1))
    for typ, m in u.servis_hel.items():
        wb.satt(UPPG, f"{SERVIS_HEL[typ]}{RAD0_UPPG}", round(m, 1))
    if u.servis_ventilgropar:
        wb.satt(UPPG, f"{GROP_SERVIS_UPPG}{RAD0_UPPG}", u.servis_ventilgropar)
    if u.inforingsgropar:
        wb.satt(ENTR, ENTR_INFORINGSGROP, u.inforingsgropar)
        wb.satt(ENTR, ENTR_MASKINGROP, u.inforingsgropar)
    for typ, m in u.servis_tillagg.items():
        wb.satt(ENTR, ENTR_TILLAGG[typ], round(m, 1))
    djup_l, djup_h = prislistor["djup"]
    for pris, m in u.djup.items():
        ref = ENTR_DJUP[0] if djup_l is not None and abs(pris - djup_l) < 0.5 else ENTR_DJUP[1]
        wb.satt(ENTR, ref, round(m, 1))
    r = ENTR_SPONT_RAD
    if u.spont_m2 > 0:
        wb.satt(ENTR, f"B{r}", f"Slagen spont båda sidor, höjd medeldjup + {la.SPONT_NEDSLAGNING_M:g} m "
                               f"(ledning_analys, schakt efter djup)")
        wb.satt(ENTR, f"C{r}", "m2")
        wb.satt(ENTR, f"D{r}", round(u.spont_m2, 1))
        wb.satt(ENTR, f"F{r}", la.kr(kost, "slagen_spont") or 0.0)
    prov = _vanligast(u.provisorier)
    if prov:
        r = ENTR_PROV_RAD
        wb.satt(ENTR, f"B{r}", f"Provisorier {100 * prov:g} % av entreprenaden (områdets fält "
                               f"{la.PROVISORIER_FALT})")
        wb.satt(ENTR, f"C{r}", "st")
        wb.satt(ENTR, f"D{r}", 1)
        wb.satt(ENTR, f"F{r}", Formel(f"(H55+H91+H132+SUM(H136:H{r - 1}))*{prov!r}"))
    trafik = _vanligast(u.trafik)
    if trafik:
        r = ENTR_TRAFIK_RAD
        wb.satt(ENTR, f"F{r}", 1)
        wb.satt(ENTR, f"G{r}", trafik)
        wb.satt(ENTR, f"I{r}", f"Trafiktillägg {100 * trafik:g} % ur områdets fält {la.TRAFIK_FALT}")

    # byggherrekostnaderna: procent ur fliken Alternativ, kända risker som fritextrad
    for rad in RIML_GRUPPER:
        wb.satt(RIML, f"D{rad}", Formel(f"C{rad}"))
        wb.satt(RIML, f"I{rad}", "Procent av entreprenaden ur fliken Alternativ (ifyllt av ledning_analys.py)")
    risk = la.KANDA_RISKER_ANDEL
    for rad, text, formel in (
            (RISK_PLANERA_RAD, f"Kända risker {100 * risk:g} % av byggherrekostnaderna (schablon)",
             f"{risk!r}*SUM(Alternativ!D23:D26)"),
            (RISK_GENOMFORA_RAD, f"Kända risker {100 * risk:g} % av byggherrekostnaderna och entreprenaden (schablon)",
             f"{risk!r}*SUM(Alternativ!D31:D35)")):
        wb.satt(RISK, f"B{rad}", text)
        wb.satt(RISK, f"D{rad}", 1)
        wb.satt(RISK, f"G{rad}", Formel(formel))
        wb.satt(RISK, f"J{rad}", 1)

    # uppgifter om utredningen
    def m_text(x):
        return f"{x:,.0f}".replace(",", " ")
    metoder = ", ".join(f"{namn} {m_text(u.metod_m[k])} m"
                        for k, namn in (("strumpa", "infodring"), ("rörspräckning", "rörspräckning"),
                                        ("omläggning", "omläggning")) if u.metod_m.get(k))
    beskr = (f"{alt_text}. Ur ledningsinventeringen (ledning_analys.py {datum}): {u.ledningar} ledningar, "
             f"{m_text(u.langd)} m" + (f" – {metoder}" if metoder else "") + "."
             + (f" {len(u.utan_pris)} ledningar saknar pris och ingår inte: {', '.join(u.utan_pris[:20])}"
                + (" …" if len(u.utan_pris) > 20 else "") + "." if u.utan_pris else ""))
    wb.satt(UPPG, "G8", f"{u.namn} – {'alt. schakt' if u.alternativ == 'schakt' else 'alt. schaktfritt'}")
    wb.satt(UPPG, "G13", alt_nr)
    wb.satt(UPPG, "G15", beskr)
    wb.spara(ut)


def _filnamn(text: str) -> str:
    t = text.translate(str.maketrans("åäöÅÄÖéÉ", "aaoAAOeE"))
    return re.sub(r"[^A-Za-z0-9]+", "_", t).strip("_") or "omrade"


def skriv(data: dict, ledningar: list[dict], kost, mall: str, ut_mapp: str, la) -> list[str]:
    """Skriver en ifylld kalkylmall per område och alternativ i ut_mapp. Returnerar filerna."""
    from datetime import datetime
    prislistor = las_prislistor(mall)
    os.makedirs(ut_mapp, exist_ok=True)
    datum = datetime.now().strftime("%Y-%m-%d")
    filer = []
    for (omr, alt), u in sorted(underlag(data, ledningar, kost, la, prislistor).items()):
        if not u.ledningar:
            continue
        nr, text = next((n, t) for a, n, t in ALTERNATIV if a == alt)
        path = os.path.join(ut_mapp, f"Kalkyl_{_filnamn(omr)}_{alt}.xlsx")
        fyll(mall, path, u, nr, text, prislistor, kost, la, datum)
        filer.append(path)
    return filer
