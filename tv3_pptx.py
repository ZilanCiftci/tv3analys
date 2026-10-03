#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
tv3_pptx.py – PowerPoint ur tv3_analys (PLAN_verktyg.md steg 4).

Fyller mallen mall.pptx (layouterna TITEL, AVDELARE, RUBRIK, RUBRIK_TEXT, RUBRIK_BILD,
RUBRIK_TABELL, STRACKA, AVSLUT med numrerade platshållare) med resultatet av analysen.

Två sätt att köra:
  python tv3_analys.py -l filer.txt --pptx            # presentationen skrivs i utdatamappen
  python tv3_pptx.py tv3_resultat [--mall mall.pptx] [--ut fil.pptx] [--topp 10]
                                                       # efteråt, ur en färdig utdatamapp (läser
                                                       # kartunderlag.json, prioritering.xlsx med
                                                       # manuella bedömningar och TV3-filerna igen)

Sidorna styrs av SIDOR i KONFIG; stryk det som inte ska med. Finns kartbild.png i utdatamappen
(ArcMap-verktyget Exportera kartbild) läggs den in som egen sida. Beroende: python-pptx.
Mallen får anpassas fritt i PowerPoint så länge layoutnamnen och platshållarnas ordning behålls.
"""
from __future__ import annotations

import argparse
import os
import sys
from collections import Counter
from datetime import datetime

# ----------------------------------------------------------------------------
# KONFIG
# ----------------------------------------------------------------------------

MALL = os.path.join(os.path.dirname(os.path.abspath(__file__)), "mall.pptx")
PPTX_FIL = "presentation.pptx"         # skrivs i utdatamappen
KARTBILD = "kartbild.png"              # från ArcMap-verktyget Exportera kartbild, i utdatamappen
TOPP_STRACKOR = 10                     # en STRACKA-sida per sträcka i toppen
TABELLRADER = 12                       # rader per tabellsida (fler → fortsättningssida)
FOTON_PER_STRACKA = 3
# Sidor i ordning. Ta bort det som inte ska med.
SIDOR = ["titel", "avdelare", "nyckeltal", "klasser", "topplista", "topptabell", "koder", "material",
         "etapper", "strackor", "driftatgarder", "svackor", "karta", "metod", "avslut"]
TITEL = "Tillståndsbedömning av avloppsledningar"
AVSLUT_PUNKTER = [
    "Manuell genomgång av klass A och B i fält/film – bedömningen förs in i Excel (Manuell bedömning)",
    "Beslut om etappindelning och metod per etapp",
    "Kostnadsbedömning av schaktsträckor (djup, spont)",
    "Driftåtgärder (rotskärning, spolning) beställs separat",
]

# Platshållarnas idx per layout (ordningen i mallen; namnen är inte stabila i PowerPoint)
PH = {
    "TITEL": {"title": 100, "undertitel": 101, "meta": 102},
    "AVDELARE": {"title": 100, "body": 101},
    "RUBRIK": {"title": 100},
    "RUBRIK_TEXT": {"title": 100, "body": 101},
    "RUBRIK_BILD": {"title": 100, "bild": 101, "body": 102},
    "RUBRIK_TABELL": {"title": 100, "tabell": 101},
    "STRACKA": {"title": 100, "oversikt": 101, "foto1": 102, "foto2": 103, "foto3": 104, "fakta": 105},
    "AVSLUT": {"title": 100, "body": 101},
}

# ----------------------------------------------------------------------------


def _tal(v: float, dec: int = 0) -> str:
    """Svenskt talformat: tusentalsmellanslag och decimalkomma."""
    s = f"{v:,.{dec}f}".replace(",", " ").replace(".", ",")
    return s


class Mall:
    """Hjälp kring python-pptx: layouter, platshållare, bilder, tabeller och nyckeltalsrutor."""

    def __init__(self, mall: str):
        from pptx import Presentation
        from pptx.util import Emu
        self.prs = Presentation(mall)
        self.Emu = Emu
        self._rensa_exempelsidor()
        self.layouter = {l.name: l for l in self.prs.slide_layouts}
        saknas = [n for n in PH if n not in self.layouter]
        if saknas:
            raise SystemExit(f"Mallen saknar layouterna {', '.join(saknas)} – använd mall.pptx från repot "
                             "eller behåll layoutnamnen när mallen anpassas.")

    def _rensa_exempelsidor(self):
        """Tar bort exempelsidorna i mallen (layouterna behålls)."""
        lst = self.prs.slides._sldIdLst
        for sld in list(lst):
            rid = sld.rId
            self.prs.part.drop_rel(rid)
            lst.remove(sld)

    # --- sidor och platshållare -------------------------------------------------
    def sida(self, layout: str):
        slide = self.prs.slides.add_slide(self.layouter[layout])
        # Sidnummer ligger som platshållare i layouten; python-pptx kopierar den inte automatiskt
        try:
            from pptx.enum.shapes import PP_PLACEHOLDER
            for ph in self.layouter[layout].placeholders:
                if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
                    slide.shapes.clone_layout_placeholder(ph)
        except Exception:
            pass
        return slide

    @staticmethod
    def ph(slide, layout: str, namn: str):
        idx = PH[layout][namn]
        for sh in slide.placeholders:
            if sh.placeholder_format.idx == idx:
                return sh
        return None

    def text(self, slide, layout: str, namn: str, rader, storlek: float | None = None):
        """Text i en platshållare. rader: str eller lista av str / (str, fet) / (str, fet, nivå)."""
        from pptx.util import Pt
        ph = self.ph(slide, layout, namn)
        if ph is None:
            return
        if isinstance(rader, str):
            rader = [rader]
        tf = ph.text_frame
        tf.clear()
        first = True
        for rad in rader:
            fet, niva = False, 0
            if isinstance(rad, tuple):
                txt = rad[0]
                fet = rad[1] if len(rad) > 1 else False
                niva = rad[2] if len(rad) > 2 else 0
            else:
                txt = rad
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.level = niva
            r = p.add_run()
            r.text = str(txt)
            if fet:
                r.font.bold = True
            if storlek:
                r.font.size = Pt(storlek)

    def bild(self, slide, layout: str, namn: str, path: str | None):
        """Bild inpassad i platshållarens ruta (behåller proportionerna, centrerad). Platshållaren
        tas bort; saknas bilden tas den bara bort så att ingen tom ruta visas."""
        ph = self.ph(slide, layout, namn)
        if ph is None:
            return
        left, top, w, h = ph.left, ph.top, ph.width, ph.height
        ph._element.getparent().remove(ph._element)
        if not path or not os.path.isfile(path):
            return
        try:
            from PIL import Image
            with Image.open(path) as im:
                bw, bh = im.size
        except Exception:
            bw, bh = 4, 3
        skala = min(w / bw, h / bh)
        nw, nh = int(bw * skala), int(bh * skala)
        slide.shapes.add_picture(path, left + (w - nw) // 2, top + (h - nh) // 2, nw, nh)

    def tabell(self, slide, layout: str, namn: str, kolumner: list[str], rader: list[list],
               bredder: list[float] | None = None, storlek: float = 12, klasskol: int | None = None):
        """Tabell i platshållarens ruta. bredder: relativa kolumnbredder. klasskol: kolumn vars
        första bokstav (A–E) färgar cellen med klassfärgen."""
        from pptx.dml.color import RGBColor
        from pptx.enum.dml import MSO_THEME_COLOR
        from pptx.util import Pt
        ph = self.ph(slide, layout, namn)
        if ph is None:
            return
        left, top, w, h = ph.left, ph.top, ph.width, ph.height
        ph._element.getparent().remove(ph._element)
        n_rad = len(rader) + 1
        radh = min(int(h / n_rad), self.Emu(457200 * 0.45))
        shape = slide.shapes.add_table(n_rad, len(kolumner), left, top, w, radh * n_rad)
        t = shape.table
        if bredder:
            tot = sum(bredder)
            for i, b in enumerate(bredder):
                t.columns[i].width = int(w * b / tot)
        for i in range(n_rad):
            t.rows[i].height = radh

        def satt(cell, txt, fet=False, vit=False):
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = "" if txt is None else str(txt)
            r.font.size = Pt(storlek)
            r.font.bold = fet
            if vit:
                r.font.color.theme_color = MSO_THEME_COLOR.BACKGROUND_1
            cell.margin_top = cell.margin_bottom = Pt(2)
            cell.margin_left = cell.margin_right = Pt(5)

        for j, k in enumerate(kolumner):
            c = t.cell(0, j)
            satt(c, k, fet=True, vit=True)
            c.fill.solid()
            c.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
        for i, rad in enumerate(rader, 1):
            for j, v in enumerate(rad):
                c = t.cell(i, j)
                satt(c, v)
                c.fill.solid()
                c.fill.fore_color.theme_color = MSO_THEME_COLOR.BACKGROUND_1 if i % 2 else MSO_THEME_COLOR.BACKGROUND_2
                if klasskol is not None and j == klasskol and v and str(v)[0] in KLASS_FARG:
                    c.fill.fore_color.rgb = RGBColor.from_string(KLASS_FARG[str(v)[0]])
                    r = c.text_frame.paragraphs[0].runs[0]
                    r.font.bold = True
                    if str(v)[0] in "AB":
                        r.font.color.rgb = RGBColor(255, 255, 255)
        return t

    def nyckeltal(self, slide, rutor: list[tuple[str, str]]):
        """Nyckeltalsrutor (värde, etikett) i tre kolumner på en RUBRIK-sida, som i mallens exempel."""
        from pptx.enum.dml import MSO_THEME_COLOR
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
        xs, ys = (0.6, 4.7, 8.8), (1.7, 4.2)
        for n, (varde, etikett) in enumerate(rutor[:6]):
            x, y = xs[n % 3], ys[n // 3]
            ruta = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(3.8), Inches(2.1))
            ruta.name = f"stat-ruta-{n + 1}"
            ruta.adjustments[0] = 0.08
            ruta.fill.solid()
            ruta.fill.fore_color.theme_color = MSO_THEME_COLOR.BACKGROUND_2
            ruta.line.fill.background()
            ruta.shadow.inherit = False
            for namn, txt, dy, hh, pt, fet, farg in (
                    ("varde", varde, 0.3, 1.0, 40, True, MSO_THEME_COLOR.TEXT_2),
                    ("etikett", etikett, 1.3, 0.6, 14, False, MSO_THEME_COLOR.TEXT_1)):
                tb = slide.shapes.add_textbox(Inches(x + 0.3), Inches(y + dy), Inches(3.2), Inches(hh))
                tb.name = f"stat-{namn}-{n + 1}"
                tf = tb.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.LEFT
                r = p.add_run()
                r.text = txt
                r.font.size = Pt(pt)
                r.font.bold = fet
                r.font.color.theme_color = farg

    def spara(self, path: str):
        self.prs.save(path)


KLASS_FARG = {"A": "D03B3B", "B": "EC835A", "C": "FAB219", "D": "0CA30C", "E": "BFBFBF"}


# ----------------------------------------------------------------------------
# Innehåll
# ----------------------------------------------------------------------------

def _period(strackor) -> str:
    d = sorted({s.datum for s in strackor if s.datum})
    if not d:
        return ""
    return d[0] if d[0] == d[-1] else f"{d[0]} – {d[-1]}"


def _strackanamn(s) -> str:
    return f"{s.startbrunn} → {s.slutbrunn}"


def bygg_presentation(strackor, etapper, diagram: dict, utdata: str, mall: str | None = None,
                      ut: str | None = None, topp: int = TOPP_STRACKOR, sidor=None, ta=None) -> str:
    """Bygger presentationen. strackor/etapper från tv3_analys (ta = modulen tv3_analys, för
    KONFIG-värden och ritfunktioner), diagram = {"klasser", "topp", "koder", "material"} → PNG.
    Returnerar sökvägen till .pptx-filen."""
    import tempfile
    ta = ta or sys.modules.get("tv3_analys") or __import__("tv3_analys")
    sidor = list(sidor or SIDOR)
    m = Mall(mall or MALL)
    ut = ut or os.path.join(utdata, PPTX_FIL)
    sorterade = ta.sorterade_strackor(strackor)
    bedomda = [s for s in strackor if s.klass != "E"]
    klasser = Counter(s.klass for s in strackor)
    langd = sum(s.langd for s in strackor)
    langd_k = {k: sum(s.langd for s in strackor if s.klass == k) for k in "ABCDE"}
    omraden = sorted({s.omrade for s in strackor if s.omrade})
    filer = sorted({os.path.splitext(os.path.basename(s.fil))[0] for s in strackor})
    agare = sorted({getattr(s, "agare", "") for s in strackor if getattr(s, "agare", "")})
    tmp = tempfile.mkdtemp(prefix="tv3_pptx_")

    def rubrik_tabell(titel, kolumner, rader, bredder, klasskol=None, storlek=12):
        """En eller flera tabellsidor (TABELLRADER rader per sida)."""
        for i in range(0, max(len(rader), 1), TABELLRADER):
            slide = m.sida("RUBRIK_TABELL")
            m.text(slide, "RUBRIK_TABELL", "title", titel + (f" ({i // TABELLRADER + 1})" if i else ""))
            m.tabell(slide, "RUBRIK_TABELL", "tabell", kolumner, rader[i:i + TABELLRADER], bredder,
                     storlek=storlek, klasskol=klasskol)

    for sida in sidor:
        if sida == "titel":
            slide = m.sida("TITEL")
            m.text(slide, "TITEL", "title", TITEL)
            m.text(slide, "TITEL", "undertitel",
                   "  ·  ".join(x for x in [", ".join(filer[:4]) + (" …" if len(filer) > 4 else ""),
                                            " / ".join(omraden[:3]), f"TV-inspektion {_period(strackor)}"] if x.strip()))
            m.text(slide, "TITEL", "meta",
                   "  ·  ".join(x for x in [", ".join(agare), "underlag för renoveringsplanering",
                                            datetime.now().strftime("%Y-%m-%d")] if x))
        elif sida == "avdelare":
            slide = m.sida("AVDELARE")
            m.text(slide, "AVDELARE", "title", "Resultat")
            m.text(slide, "AVDELARE", "body",
                   f"{len(bedomda)} sträckor och {_tal(langd)} m bedömda enligt poängmodellen (P93)")
        elif sida == "nyckeltal":
            slide = m.sida("RUBRIK")
            m.text(slide, "RUBRIK", "title", "Nyckeltal")
            andel_a = langd_k["A"] / langd * 100 if langd else 0
            m.nyckeltal(slide, [
                (str(len(strackor)), "sträckor inspekterade"),
                (f"{_tal(langd)} m", "inspekterad längd"),
                (str(klasser.get("A", 0)), f"sträckor i klass A – {ta.KLASS_TEXT['A'].split('– ')[-1]}"),
                (f"{_tal(langd_k['A'])} m", f"längd i klass A ({andel_a:.0f} %)"),
                (str(sum(len(s.skador()) for s in strackor)), "registrerade skador"),
                (str(sum(1 for s in strackor if s.avbruten)), "avbrutna inspektioner"),
            ])
        elif sida == "klasser":
            slide = m.sida("RUBRIK_BILD")
            m.text(slide, "RUBRIK_BILD", "title", "Inspekterad längd per prioritetsklass")
            m.bild(slide, "RUBRIK_BILD", "bild", diagram.get("klasser"))
            rader = []
            for k in "ABCD":
                rader.append((ta.KLASS_TEXT[k], True))
                rader.append(f"{klasser.get(k, 0)} sträckor, {_tal(langd_k[k])} m")
            m.text(slide, "RUBRIK_BILD", "body", rader, storlek=13)
        elif sida == "topplista":
            slide = m.sida("RUBRIK_BILD")
            m.text(slide, "RUBRIK_BILD", "title", f"De {min(topp, len(sorterade))} mest kritiska sträckorna")
            m.bild(slide, "RUBRIK_BILD", "bild", diagram.get("topp"))
            m.text(slide, "RUBRIK_BILD", "body", [
                "Konstruktionsindex i poäng per 100 m. Driftskador (rötter, sediment, inläckage) ingår inte i "
                "klassen utan redovisas som driftåtgärd.",
                "",
                (f"Klass A: grad 4 på {'/'.join(sorted(ta.GRAD4_KODER_A))} eller index ≥ {ta.TROSKEL_A:g}", False),
                (f"Klass B: grad 3 eller index ≥ {ta.TROSKEL_B:g}", False),
            ], storlek=13)
        elif sida == "topptabell":
            rader = []
            for rang, s in enumerate(sorterade[:topp], 1):
                rader.append([rang, _strackanamn(s), s.material, s.dimension, f"{s.langd:.0f}",
                              f"{s.index('K'):.0f}", s.sammanfattning_skador()[:60]])
            rubrik_tabell(f"Topplista – de {len(rader)} mest kritiska",
                          ["Rang", "Sträcka", "Material", "Dim", "Längd (m)", "Index", "Skador"],
                          rader, [0.8, 3.2, 1.3, 0.9, 1.2, 1.0, 3.7], storlek=11)
        elif sida == "koder":
            slide = m.sida("RUBRIK_BILD")
            m.text(slide, "RUBRIK_BILD", "title", "Observationer per skadekod")
            m.bild(slide, "RUBRIK_BILD", "bild", diagram.get("koder"))
            c = Counter(o.kod for s in strackor for o in s.skador())
            rader = []
            for kod, n in c.most_common(6):
                rader.append((f"{ta.KODER.get(kod, (kod, ''))[0]} ({kod})", True))
                g4 = sum(1 for s in strackor for o in s.skador() if o.kod == kod and o.grad == 4)
                rader.append(f"{n} st" + (f", varav {g4} grad 4" if g4 else ""))
            m.text(slide, "RUBRIK_BILD", "body", rader, storlek=12)
        elif sida == "material":
            slide = m.sida("RUBRIK_BILD")
            m.text(slide, "RUBRIK_BILD", "title", "Prioritetsklass per material")
            m.bild(slide, "RUBRIK_BILD", "bild", diagram.get("material"))
            grp = Counter()
            ab = Counter()
            for s in strackor:
                grp[s.material_grupp or "Okänt"] += s.langd
                if s.klass in "AB":
                    ab[s.material_grupp or "Okänt"] += s.langd
            rader = []
            for mat, L in grp.most_common(6):
                rader.append((mat, True))
                rader.append(f"{_tal(L)} m, {ab[mat] / L * 100 if L else 0:.0f} % i klass A+B")
            m.text(slide, "RUBRIK_BILD", "body", rader, storlek=12)
        elif sida == "etapper" and etapper:
            rader = []
            for e in etapper:
                st = e["strackor"]
                rader.append([e["nr"], e["metod"], min(s.gallande_klass for s in st), len(st),
                              _tal(e["langd_m"]), len(e["framschaktade"]),
                              _tal(e["kostnad"]["summa"]) if e["metod"] == "strumpa" else "ej beräknad"])
            tot = sum(e["kostnad"]["summa"] for e in etapper)
            rader.append(["", "Summa", "", sum(len(e["strackor"]) for e in etapper),
                          _tal(sum(e["langd_m"] for e in etapper)),
                          sum(len(e["framschaktade"]) for e in etapper), _tal(tot)])
            rubrik_tabell("Etapper och kostnad (schablon)",
                          ["Etapp", "Metod", "Klass", "Sträckor", "Längd (m)", "Brunnar att schakta fram", "Kostnad (kr)"],
                          rader, [0.8, 1.2, 0.8, 1.0, 1.2, 2.2, 1.8], klasskol=2, storlek=11)
        elif sida == "strackor":
            for rang, s in enumerate(sorterade[:topp], 1):
                slide = m.sida("STRACKA")
                m.text(slide, "STRACKA", "title", f"{rang}. {_strackanamn(s)}  ·  Klass {s.gallande_klass}")
                ov = os.path.join(tmp, f"oversikt_{rang}.png")
                try:
                    ta.rita_schema(s, ov)
                except Exception:
                    ov = None
                m.bild(slide, "STRACKA", "oversikt", ov)
                # Foton: de allvarligaste skadorna först, en bild per observation
                obs = sorted((o for o in s.observationer if any(p for n, p in o.bilder)),
                             key=lambda o: (-(o.grad or 0) if o.raknas else 1, o.lage))
                foton, sedda = [], set()
                for o in obs:
                    for n, p in o.bilder:
                        if p and p not in sedda:
                            foton.append(p)
                            sedda.add(p)
                            break
                    if len(foton) >= FOTON_PER_STRACKA:
                        break
                for i in range(FOTON_PER_STRACKA):
                    m.bild(slide, "STRACKA", f"foto{i + 1}", foton[i] if i < len(foton) else None)
                pa = s.profil_analys
                fakta = [
                    (f"{s.material} {s.dimension} mm, {_tal(s.langd, 1)} m, {s.ledningstyp.lower()}", True),
                    f"Konstruktionsindex {s.index('K'):.0f} p/100 m, maxgrad {s.maxgrad('K') or '–'}",
                    f"Skador: {s.sammanfattning_skador() or 'inga'}",
                    f"Anslutningar: {s.antal_anslutningar}",
                ]
                if pa and pa["svackdjup"] is not None and pa["svackdjup"] >= 0.02:
                    fakta.append(f"Svacka {pa['svackdjup'] * 100:.0f} cm över {pa['svacklangd']:.0f} m")
                if s.avbruten:
                    fakta.append("Avbruten inspektion")
                if s.relinad:
                    fakta.append("Relinad")
                if s.driftatgard:
                    fakta.append(f"Driftåtgärd: {s.driftatgard}")
                if s.etapp:
                    k = f", {_tal(s.kostnad['summa'])} kr" if s.kostnad else ""
                    fakta.append(f"Etapp {s.etapp}: {s.metod}{k}")
                if s.manuell_bedomning:
                    fakta.append(f"Manuell bedömning: {s.manuell_bedomning}"
                                 + (f" – {s.kommentar}" if s.kommentar else ""))
                fakta.append(f"{s.fil}, sträcka {s.nr}, {s.datum}")
                m.text(slide, "STRACKA", "fakta", fakta, storlek=12)
        elif sida == "driftatgarder":
            drift = sorted((s for s in strackor if s.driftatgard), key=lambda s: -s.index("D"))
            if drift:
                rader = [[_strackanamn(s), s.driftatgard, f"{s.index('D'):.0f}",
                          ", ".join(f"{n}×{k}{g}" for (k, g), n in
                                    Counter((o.kod, o.grad) for o in s.skador("D")).most_common(3)),
                          s.klass] for s in drift[:TABELLRADER * 2]]
                rubrik_tabell(f"Driftåtgärder ({len(drift)} sträckor)",
                              ["Sträcka", "Åtgärd", "Driftindex", "Driftskador", "Klass"],
                              rader, [3.0, 3.2, 1.2, 3.0, 0.8], klasskol=4, storlek=11)
        elif sida == "svackor":
            sv = [(s, s.profil_analys) for s in strackor]
            sv = sorted((x for x in sv if x[1] and x[1]["svackdjup"] is not None and x[1]["svackdjup"] >= 0.02),
                        key=lambda x: -x[1]["svackdjup"])
            if sv:
                rader = [[_strackanamn(s), f"{pa['svackdjup'] * 100:.0f}",
                          f"{s.svacka_andel * 100:.0f} %" if s.svacka_andel is not None else "",
                          f"{pa['svacklangd']:.0f}", f"{pa['bakfall']:.0f}" if pa["bakfall"] >= 1 else "",
                          f"{s.material} {s.dimension}", "osäker" if pa["osaker"] else ""]
                         for s, pa in sv[:TABELLRADER * 2]]
                rubrik_tabell(f"Svackor och bakfall ({len(sv)} sträckor med svacka ≥ 2 cm)",
                              ["Sträcka", "Svackdjup (cm)", "Andel av diam.", "Stående vatten (m)",
                               "Bakfall (m)", "Material/dim", "Profil"],
                              rader, [3.0, 1.3, 1.3, 1.5, 1.1, 2.0, 1.0], storlek=11)
        elif sida == "karta":
            kb = os.path.join(utdata, KARTBILD)
            if os.path.isfile(kb):
                slide = m.sida("RUBRIK_BILD")
                m.text(slide, "RUBRIK_BILD", "title", "Bedömda ledningar i kartan")
                m.bild(slide, "RUBRIK_BILD", "bild", kb)
                m.text(slide, "RUBRIK_BILD", "body", [
                    "Färg efter prioritetsklass", "Streckad linje = maskinell bedömning",
                    "Heldragen linje = manuell bedömning", "",
                    f"{klasser.get('A', 0)} sträckor klass A, {klasser.get('B', 0)} klass B"], storlek=13)
        elif sida == "metod":
            slide = m.sida("RUBRIK_TEXT")
            m.text(slide, "RUBRIK_TEXT", "title", "Metod – poängmodell enligt P93")
            gp = ", ".join(f"grad {g} = {p} p" for g, p in sorted(ta.GRADPOANG.items()))
            kf = ", ".join(f"{k} {v:g}" for k, v in (ta.KODFAKTOR or {}).items())
            m.text(slide, "RUBRIK_TEXT", "body", [
                ("Poäng per observation", True), (gp, False, 1),
                (f"Driftskador × {ta.DRIFTFAKTOR:g}; löpande skador × längd / {ta.LOPANDE_ENHET_M:g} m, högst {ta.LOPANDE_TAK:g}", False, 1),
                (f"Kodvikter: {kf}; ytskada grad 4 = 1,0; cirkulära sprickor × {ta.ATTRIBUTFAKTOR.get(('SPR', 'CIRK'), 1):g}", False, 1),
                ("Konstruktionsindex", True),
                (f"poäng / max(längd, {ta.MINLANGD:g} m) × 100 – bara konstruktionsskador styr klassen", False, 1),
                ("Prioritetsklass", True),
                (f"A: grad 4 på {'/'.join(sorted(ta.GRAD4_KODER_A))} eller index ≥ {ta.TROSKEL_A:g}   ·   "
                 f"B: grad 3 eller index ≥ {ta.TROSKEL_B:g}   ·   C: övriga skador   ·   D: inga skador", False, 1),
                ("Åtgärdspaket", True),
                ("Klass A/B → strumpa (schakt bara manuellt); etapper av sammanhängande sträckor; "
                 "framschaktning per brunn när bara tillsynsbrunnar finns; schablonkostnader", False, 1),
            ], storlek=14)
        elif sida == "avslut":
            slide = m.sida("AVSLUT")
            m.text(slide, "AVSLUT", "title", "Nästa steg")
            m.text(slide, "AVSLUT", "body", AVSLUT_PUNKTER, storlek=18)

    m.spara(ut)
    import shutil
    shutil.rmtree(tmp, ignore_errors=True)
    return ut


# ----------------------------------------------------------------------------
# Fristående körning ur en utdatamapp
# ----------------------------------------------------------------------------

def main(argv=None):
    ap = argparse.ArgumentParser(description="PowerPoint ur en utdatamapp från tv3_analys")
    ap.add_argument("utdata", help="utdatamapp med kartunderlag.json (och prioritering.xlsx)")
    ap.add_argument("--mall", default=MALL, help="pptx-mall (standard: mall.pptx bredvid skriptet)")
    ap.add_argument("--ut", default=None, help=f"utfil (standard: <utdata>/{PPTX_FIL})")
    ap.add_argument("--topp", type=int, default=TOPP_STRACKOR, help="antal sträckor med egen sida")
    ap.add_argument("--kostnader", default=None, help="kostnadsfil (standard: kostnader.csv bredvid skriptet)")
    a = ap.parse_args(argv)

    import json
    import tempfile
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import tv3_analys as ta

    kart = os.path.join(a.utdata, ta.KARTUNDERLAG_FIL)
    if not os.path.isfile(kart):
        sys.exit(f"Hittar inte {kart} – kör tv3_analys.py först (med --karta ja).")
    with open(kart, encoding="utf-8") as f:
        data = json.load(f)
    filer = sorted({p["tv3_fil"] for p in data.get("strackor", []) if p.get("tv3_fil")})
    strackor = []
    for p in filer:
        if not os.path.isfile(p):
            print(f"  SAKNAS  {p}")
            continue
        strackor += ta.las_tv3(p)
    if not strackor:
        sys.exit("Inga TV3-filer kunde läsas – sökvägarna i kartunderlag.json pekar på en annan dator?")
    # Ersättningslittera finns inte i JSON-filen: använd de rättade namnen därifrån (fil + nr)
    rattade = {(os.path.basename(p["tv3_fil"]), p["nr"]): p for p in data.get("strackor", []) if p.get("littera_rattat")}
    for s in strackor:
        p = rattade.get((os.path.basename(s.tv3_sokvag), s.nr))
        if p:
            s.startbrunn, s.slutbrunn, s.littera_rattat = p["startbrunn"], p["slutbrunn"], p["littera_rattat"]
    media = data.get("mediakataloger", [])
    bild = data.get("bildkataloger", [])
    ta.koppla_media(strackor, media, bild)
    xl = os.path.join(a.utdata, "prioritering.xlsx")
    if os.path.isfile(xl):
        try:
            poster, fram = ta.las_manuella(xl)
            n = ta.koppla_manuella(strackor, poster)
            print(f"  {n} manuella bedömningar ur {xl}")
        except Exception as e:  # noqa: BLE001
            print(f"  kunde inte läsa manuella bedömningar ur {xl}: {e}")
            fram = []
    else:
        fram = []
    kostfil = a.kostnader or os.path.join(os.path.dirname(os.path.abspath(__file__)), ta.KOSTNADSFIL)
    kostnader = ta.las_kostnader(kostfil) if os.path.isfile(kostfil) else []
    etapper = ta.planera_atgarder(strackor, kostnader, fram)
    tmp = tempfile.mkdtemp(prefix="tv3_diagram_")
    diagram = ta.rita_diagram(strackor, tmp, a.topp)
    ut = bygg_presentation(strackor, etapper, diagram, a.utdata, a.mall, a.ut, a.topp, ta=ta)
    import shutil
    shutil.rmtree(tmp, ignore_errors=True)
    print(f"Presentation skriven: {ut}")


if __name__ == "__main__":
    main()
